"""Indexación de contenido de archivos + OCR — FileMeshy Worker (Ronda 6).

Extrae texto real de un archivo (texto plano, PDF -- con fallback a OCR si
es un PDF escaneado --, Office moderno, o directamente OCR sobre una
imagen) y lo guarda en `file_content_index` (ver `app.files.models.
FileContentIndex`) para que `CRUDFile.search` pueda matchear por contenido,
no solo por nombre.

Disparada (encolada, nunca esperada de forma síncrona por el request que la
originó) desde `app/files/router.py::_enqueue_content_indexing`, en los
cuatro puntos donde el contenido ACTUAL de un archivo puede cambiar: subida
inicial, versión nueva, restauración de una versión vieja, y reemplazo de
contenido vía share con permission='edit'.

Cualquier tipo mime fuera de lo soportado (video, audio, legacy
.doc/.xls/.ppt, binarios genéricos) queda en `extraction_status='unsupported'`
sin intentar nada -- limitación conocida de esta ronda, no bloqueante (ver
PENDIENTES.md)."""

import io
import logging
import os
from typing import Optional, Tuple
from uuid import UUID

from celery.exceptions import SoftTimeLimitExceeded

from app.worker.celery_app import celery_app
from app.worker.db import get_worker_db

logger = logging.getLogger(__name__)

# Tope de caracteres persistidos en `content_text` -- el propósito de esta
# tabla es habilitar búsqueda full-text, no archivar el documento completo.
# Un PDF/Office gigante igual se indexa, solo se trunca lo que se guarda.
MAX_CONTENT_CHARS = 100_000

# Tope de páginas a renderizar+OCRear de un PDF escaneado -- sin esto, un PDF
# de cientos de páginas cuelga el worker por minutos en una sola tarea (Celery
# no tiene otra tarea corriendo en paralelo en la misma cola mientras tanto).
MAX_PDF_OCR_PAGES = 20

# Tope de páginas de las que se extrae el texto EMBEBIDO de un PDF (pypdf).
# Independiente del tope de OCR: sin esto, un PDF de miles de páginas hace
# iterar/parsear todo el documento en una sola tarea aunque solo se persistan
# MAX_CONTENT_CHARS caracteres. Configurable por env (mismo estilo que
# av_scan.py). Además se corta apenas se acumulan MAX_CONTENT_CHARS.
MAX_PDF_TEXT_PAGES = int(os.getenv("MAX_PDF_TEXT_PAGES", "500"))

# Tope de intentos por (archivo, encolado): el contador `attempts` de
# `file_content_index` se incrementa y COMMITEA al inicio de cada ejecución,
# antes de leer nada -- así un archivo que mata al worker (OOM, SIGKILL) no
# se re-procesa para siempre por redelivery (acks_late). Se resetea a 0 al
# éxito y en cada nuevo encolado (`_enqueue_content_indexing`).
MAX_INDEX_ATTEMPTS = int(os.getenv("MAX_INDEX_ATTEMPTS", "3"))

# Tamaño máximo de archivo a indexar: `_read_private_object` lee el objeto
# completo a memoria. Por defecto reutiliza el tope de av_scan.py
# (AV_SCAN_INDEXING_MAX_SIZE_BYTES, 512MB) para que ambos caminos coincidan.
INDEXING_MAX_SIZE_BYTES = int(
    os.getenv("INDEXING_MAX_SIZE_BYTES", os.getenv("AV_SCAN_INDEXING_MAX_SIZE_BYTES", str(512 * 1024 * 1024)))
)

# Timeouts propios (segundos) para las herramientas externas (Tesseract /
# poppler): no dependen solo del soft limit de Celery.
OCR_IMAGE_TIMEOUT_SECONDS = int(os.getenv("OCR_IMAGE_TIMEOUT_SECONDS", "60"))
PDF_RENDER_TIMEOUT_SECONDS = int(os.getenv("PDF_RENDER_TIMEOUT_SECONDS", "90"))

# Timeout total de la tarea (segundos) -- OCR puede ser lento, pero tiene que
# tener un techo. `SoftTimeLimitExceeded` se captura como fallo de extracción
# (ver `index_file_content`), nunca deja el worker colgado indefinidamente.
TASK_TIME_LIMIT_SECONDS = 180
TASK_SOFT_TIME_LIMIT_SECONDS = 120

PLAIN_TEXT_MIME_TYPES = {
    "text/plain",
    "text/markdown",
    "text/csv",
    "text/html",
    "text/xml",
    "application/json",
}

# Mismo set que `INLINE_SAFE_MIME_TYPES` de app/files/router.py para el
# subconjunto de imágenes (no video/audio -- esos son 'unsupported' acá).
IMAGE_MIME_TYPES = {"image/jpeg", "image/png", "image/gif", "image/webp"}

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

# Idiomas de Tesseract -- español primero (el producto y sus usuarios son
# hispanohablantes, ver MEMORY.md/CLAUDE.md), inglés como red de seguridad
# para documentos mixtos/en inglés. Requiere el paquete de idioma
# `tesseract-ocr-spa` instalado en la imagen (ver Dockerfile).
TESSERACT_LANGS = "spa+eng"


def _read_private_object(storage_key: str) -> bytes:
    """Lee un objeto del bucket privado de MinIO -- mismo cliente boto3
    interno que usa el resto de la app (`app.files.service._get_client`),
    reimplementado acá en vez de importar `app.files.router` (que carga
    dependencias de FastAPI/HTTPException que no aplican en un proceso
    worker). Ver `app.files.router._read_private_object` para el mecanismo
    gemelo del lado API."""
    from app.core.config import settings
    from app.files.service import _get_client

    client = _get_client()
    if not client:
        raise RuntimeError("El servicio de almacenamiento (MinIO) no está disponible")
    obj = client.get_object(Bucket=settings.MINIO_PRIVATE_BUCKET, Key=storage_key)
    return obj["Body"].read()


def _decode_text(content: bytes) -> str:
    """Decodifica texto plano -- UTF-8 primero (caso común), `latin-1` como
    fallback (nunca falla: `latin-1` mapea CUALQUIER byte a un carácter)."""
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1", errors="replace")


def _extract_pdf_text(content: bytes) -> str:
    """Texto embebido de un PDF (no escaneado) vía `pypdf`. Página que falle
    individualmente se salta -- un PDF con una página corrupta no debería
    perder el texto de las demás."""
    import pypdf

    reader = pypdf.PdfReader(io.BytesIO(content))
    parts = []
    total_chars = 0
    for page_number, page in enumerate(reader.pages):
        # Tope de páginas y de caracteres acumulados (ya no se persistiría más).
        if page_number >= MAX_PDF_TEXT_PAGES or total_chars >= MAX_CONTENT_CHARS:
            break
        try:
            page_text = page.extract_text() or ""
        except SoftTimeLimitExceeded:
            # No tragarla: si no, la señal no se repite y el hard limit mata
            # el proceso.
            raise
        except Exception:
            continue
        parts.append(page_text)
        total_chars += len(page_text)
    return "\n".join(parts)


def _ocr_pdf(content: bytes) -> str:
    """Fallback de OCR para un PDF escaneado (sin texto embebido, solo
    imágenes) -- renderiza páginas a imágenes (`pdf2image`, requiere
    `poppler-utils` en el sistema, ver Dockerfile) y corre Tesseract sobre
    cada una. Limitado a `MAX_PDF_OCR_PAGES` para no colgar el worker con
    PDFs enormes."""
    from pdf2image import convert_from_bytes
    import pytesseract

    images = convert_from_bytes(
        content, first_page=1, last_page=MAX_PDF_OCR_PAGES, timeout=PDF_RENDER_TIMEOUT_SECONDS
    )
    parts = []
    for image in images:
        try:
            parts.append(pytesseract.image_to_string(image, lang=TESSERACT_LANGS, timeout=OCR_IMAGE_TIMEOUT_SECONDS))
        except SoftTimeLimitExceeded:
            raise
        except Exception:
            continue
    return "\n".join(parts)


def _extract_docx_text(content: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(content))
    return "\n".join(p.text for p in document.paragraphs)


def _extract_xlsx_text(content: bytes) -> str:
    """`data_only=True` para leer VALORES calculados, no fórmulas crudas
    (pedido explícito del diseño: "no fórmulas"). `read_only=True` evita
    cargar el workbook completo en memoria de más para archivos grandes."""
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    parts = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    parts.append(str(cell))
    workbook.close()
    return "\n".join(parts)


def _extract_pptx_text(content: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(content))
    parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
    return "\n".join(parts)


def _ocr_image(content: bytes) -> str:
    """OCR directo sobre una imagen (mismos tipos que ya se aceptan para
    preview inline, ver `INLINE_SAFE_MIME_TYPES` en el router) -- Pillow para
    abrir/normalizar el formato antes de pasarla a Tesseract."""
    from PIL import Image
    import pytesseract

    image = Image.open(io.BytesIO(content))
    return pytesseract.image_to_string(image, lang=TESSERACT_LANGS, timeout=OCR_IMAGE_TIMEOUT_SECONDS)


def _extract_text_for_mime(content: bytes, mime_type: str) -> Tuple[str, str]:
    """Devuelve `(texto_extraido, extraction_status)`. `extraction_status`
    es `'done'` o `'unsupported'` -- cualquier excepción durante la
    extracción se propaga al caller (`index_file_content`), que la traduce a
    `'failed'`."""
    if mime_type in PLAIN_TEXT_MIME_TYPES:
        return _decode_text(content), "done"

    if mime_type == "application/pdf":
        text = _extract_pdf_text(content)
        # Texto vacío o casi vacío -> probablemente un PDF escaneado (solo
        # imágenes, sin capa de texto) -- fallback a OCR. Umbral bajo (20
        # caracteres) a propósito: un PDF con texto real casi siempre trae
        # bastante más que eso incluso en una sola página.
        if len(text.strip()) < 20:
            ocr_text = _ocr_pdf(content)
            if ocr_text.strip():
                return ocr_text, "done"
        return text, "done"

    if mime_type == DOCX_MIME:
        return _extract_docx_text(content), "done"
    if mime_type == XLSX_MIME:
        return _extract_xlsx_text(content), "done"
    if mime_type == PPTX_MIME:
        return _extract_pptx_text(content), "done"

    if mime_type in IMAGE_MIME_TYPES:
        return _ocr_image(content), "done"

    # Video, audio, legacy Office (.doc/.xls/.ppt), binarios genéricos --
    # limitación conocida de esta ronda, documentada en PENDIENTES.md.
    return "", "unsupported"


@celery_app.task(
    bind=True,
    max_retries=0,
    # Solo esta tarea: si el worker muere a mitad (OOM/SIGKILL/hard limit) NO se
    # reencola indefinidamente (la config global es reject_on_worker_lost=True).
    reject_on_worker_lost=False,
    name="app.worker.tasks.indexing.index_file_content",
    queue="indexing",
    time_limit=TASK_TIME_LIMIT_SECONDS,
    soft_time_limit=TASK_SOFT_TIME_LIMIT_SECONDS,
)
def index_file_content(self, tenant_schema: Optional[str], file_id: str):
    """Extrae y guarda el texto indexable de `file_id`. Nunca revienta el
    worker -- cualquier excepción (incluido timeout de Celery) se captura y
    se persiste como `extraction_status='failed'` con un mensaje corto (no
    el stack trace completo) en `error_message`.

    `max_retries=0` a propósito: un reintento automático de OCR sobre el
    mismo documento que ya falló (ej. corrupto, formato inesperado) no va a
    dar un resultado distinto -- si algún día se agrega reintento real,
    debería ser para errores transitorios de infraestructura (MinIO caído),
    no para fallos de parseo del documento en sí.

    NO hace falta acá ningún import defensivo de "todos los modelos" para
    que `configure_mappers()` no reviente con las relaciones ORM que
    resuelven "User"/"Role"/etc. como string -- eso está centralizado en
    `app.worker.celery_app._register_all_orm_models()` (se ejecuta solo, a
    nivel de módulo y vía la señal `worker_process_init`, ANTES de que
    cualquier tarea corra) desde que se encontró en producción que
    duplicar ese workaround ad-hoc en cada tarea (como tenía antes esta
    misma función) no alcanza: si otra tarea sin el workaround corre
    PRIMERO en el mismo proceso worker, el registry queda roto para
    SIEMPRE en ese proceso, y ni siquiera el import correcto de acá
    (que corre DESPUÉS) lo destraba. Ver el docstring completo de esa
    función en celery_app.py para el bug real que motivó centralizarlo."""
    from app.files.repository import file_content_index_repo, file_repo

    with get_worker_db(schema_name=tenant_schema) as db:
        file_obj = file_repo.get(db, UUID(str(file_id)))
        if file_obj is None:
            logger.warning(f"index_file_content: archivo {file_id} no existe (¿borrado antes de indexar?)")
            return

        # Contador de intentos: se incrementa y commitea ANTES de leer nada.
        previous_attempts = file_content_index_repo.register_attempt(db, file_obj.id)
        if previous_attempts >= MAX_INDEX_ATTEMPTS:
            logger.error(
                f"index_file_content: {file_obj.id} superó {MAX_INDEX_ATTEMPTS} intentos "
                "(posible redelivery por caída del worker) -- se marca failed sin procesar"
            )
            _mark_failed(
                db, file_obj.id,
                f"Se superó el máximo de {MAX_INDEX_ATTEMPTS} intentos de indexación "
                "(el procesamiento anterior no terminó)",
            )
            return

        # Tope de tamaño ANTES de leer el objeto completo a memoria.
        if file_obj.size_bytes is not None and file_obj.size_bytes > INDEXING_MAX_SIZE_BYTES:
            file_content_index_repo.upsert(
                db,
                file_obj.id,
                content_text=None,
                extraction_status="unsupported",
                error_message=(
                    f"Archivo demasiado grande para indexar por contenido "
                    f"({file_obj.size_bytes} bytes, máximo {INDEXING_MAX_SIZE_BYTES})"
                ),
            )
            return

        try:
            content = _read_private_object(file_obj.storage_key)
            text, extraction_status = _extract_text_for_mime(content, file_obj.mime_type)
            content_text = text[:MAX_CONTENT_CHARS] if extraction_status == "done" else None
            file_content_index_repo.upsert(
                db,
                file_obj.id,
                content_text=content_text,
                extraction_status=extraction_status,
                attempts=0,  # éxito: resetea el contador
            )
            logger.info(
                f"index_file_content: {file_obj.id} -> {extraction_status} "
                f"({len(content_text) if content_text else 0} chars)"
            )
        except SoftTimeLimitExceeded:
            logger.error(f"index_file_content: soft time limit indexando {file_id} (tenant={tenant_schema or 'public'})")
            _mark_failed(
                db, file_obj.id,
                f"La extracción superó el tiempo máximo ({TASK_SOFT_TIME_LIMIT_SECONDS}s)",
            )
        except Exception as e:
            # Nunca el stack trace completo en `error_message` (queda
            # expuesto vía GET /files/{id}/index-status más adelante si se
            # decide mostrarlo) -- solo el mensaje, igual que el resto de
            # los manejos de error de este proyecto (ver
            # `convert_office_to_pdf` en app/files/service.py).
            logger.exception(f"index_file_content: fallo indexando {file_id} (tenant={tenant_schema or 'public'})")
            _mark_failed(db, file_obj.id, str(e))


def _mark_failed(db, file_id: UUID, message: str) -> None:
    """Persiste `extraction_status='failed'` con un mensaje corto; nunca
    lanza (ni siquiera si la propia escritura falla)."""
    from app.files.repository import file_content_index_repo

    try:
        file_content_index_repo.upsert(
            db,
            file_id,
            content_text=None,
            extraction_status="failed",
            error_message=message[:500],
        )
    except Exception:
        logger.exception(f"index_file_content: no se pudo ni siquiera registrar el fallo de {file_id}")
